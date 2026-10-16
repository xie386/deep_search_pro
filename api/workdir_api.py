"""06 · 工作目录接口层：前端文件管理器的服务端。

定案见 `Verdict/06-…md`（裁定 1 只做外部真实目录、裁定 2 默认条目「主页」、裁定 3 保护名单 fail-closed、
裁定 5 画像目录专属边界），方案见规划文档 §06 ③ E1。

**硬规矩**（与 Agent 侧写工具同一套门禁，两处口径不许漂移）：
  · 一切路径先 `workdir_guard.resolve_in_workdir`（拒 `..` / 绝对路径逃逸 / UNC），
    再按动作分档过 `assert_writable`（写）或 `check_readable`（读）；
  · 凭据类（*key*/*token*/*secret*/*credential*/*password* 与点开头目录）**连读都拒绝**，
    列表里只标记不展示内容；
  · `agents_docs/` 画像目录不在工作目录里，读写都拒；
  · **不做上传**；删除一律进 `.回收站/`（软删，可还原）；
  · 所有变更写 `data/audit/workdir.jsonl`（只记路径与体积，**不留文件内容**）。
"""
from __future__ import annotations

import os
import time
from pathlib import Path

from fastapi import HTTPException
from fastapi.responses import FileResponse
from pydantic import BaseModel

from tools import workdir_guard as guard
from tools import workdir_store as store
from utils import audit
from utils.settings import get_int

_TEXT_EXT = {".md", ".markdown", ".txt", ".json", ".jsonl", ".csv", ".tsv", ".log", ".yml", ".yaml",
             ".py", ".js", ".ts", ".html", ".htm", ".css", ".sh", ".bat", ".ini", ".cfg", ".conf", ".xml"}
_EDIT_EXT = _TEXT_EXT | {".srt", ".vtt"}          # 可编辑（前端 textarea）
_MAX_READ = 2 * 1024 * 1024                        # 预览/读取单次上限（读 2MB 文本）
_DL_EXT = {".pdf"}                                 # 仅下载


def _aid(token: str) -> int:
    from api.account import get_session
    return get_session(token)["account_id"]


def _uname(token: str) -> str:
    from api.account import get_session
    return get_session(token)["username"]


def _max_write() -> int:
    return max(1, get_int("WORKDIR_MAX_WRITE_MB", 1)) * 1024 * 1024


def _ttl_days() -> int:
    return max(1, get_int("WORKDIR_TRASH_TTL_DAYS", 7))


def _ctx(token: str):
    aid, un = _aid(token), _uname(token)
    wd, entry = store.active(aid, un)
    return aid, un, Path(wd), entry


def _audit(action: str, path: str, ok_: bool, detail: str = "", aid=None, wd=None) -> None:
    try:
        audit.write_event("workdir", action, account_id=aid, path=path, ok=bool(ok_),
                          detail=(detail or "")[:300], workdir=str(wd or ""))
    except Exception:                              # noqa: BLE001 —— 审计失败不能阻断用户操作
        pass


def _human(n: int) -> str:
    for unit, div in (("MB", 1024 * 1024), ("KB", 1024)):
        if n >= div:
            return "%.1f%s" % (n / div, unit)
    return "%dB" % n


def _rel(p: Path, wd: Path) -> str:
    try:
        return p.resolve().relative_to(wd.resolve()).as_posix()
    except ValueError:
        return p.name


# ================================================================ 模型
class AddReq(BaseModel):
    name: str = ""
    path: str
    confirm: bool = False


class IdReq(BaseModel):
    id: int


class PathReq(BaseModel):
    path: str


class CreateReq(BaseModel):
    path: str
    content: str = ""


class WriteReq(BaseModel):
    path: str
    content: str
    mtime: float | None = None


class RenameReq(BaseModel):
    path: str
    target: str


class TrashReq(BaseModel):
    action: str = "list"          # list | restore | purge
    name: str | None = None
    target: str | None = None     # restore 时可指定还原到哪个相对路径


# ================================================================ 工作目录条目
def list_(token: str) -> dict:
    aid, un, wd, entry = _ctx(token)
    items = store.list_dirs(aid, un)
    return {"ok": True, "items": items, "active_id": (entry or {}).get("id"),
            "active_path": str(wd)}


def switch_(token: str, req: IdReq) -> dict:
    aid, un, wd, _ = _ctx(token)
    r = store.switch(aid, req.id)
    _audit("switch", r.get("path", ""), bool(r.get("ok")), r.get("error", ""), aid, wd)
    if not r.get("ok"):
        raise HTTPException(400, r.get("error", "切换失败"))
    return r


RISKS = [
    "路径守卫的根目录会变成这个目录：Agent 的写工具与 CLI 面板只能在这个目录**内部**读写（内部路径规则不变）。",
    "不要在系统目录（如 C:\\Windows）或含系统文件的位置创建工作目录，避免误操作。",
    "外部目录**不在账号注销的清理链里**：删除账号时这个目录的内容不会被自动清掉，需要你自己处理。",
]


def add_(token: str, req: AddReq) -> dict:
    aid, un, wd, _ = _ctx(token)
    if not req.confirm:
        return {"ok": False, "need_confirm": True, "risks": RISKS}
    r = store.add_external(aid, req.name, req.path)
    _audit("add", r.get("path", req.path), bool(r.get("ok")), r.get("error", ""), aid, wd)
    if not r.get("ok"):
        raise HTTPException(400, r.get("error", "新增失败"))
    return r


def remove_(token: str, req: IdReq) -> dict:
    aid, un, wd, _ = _ctx(token)
    r = store.remove(aid, req.id)
    _audit("remove", "id=%s" % req.id, bool(r.get("ok")), r.get("error", ""), aid, wd)
    if not r.get("ok"):
        raise HTTPException(400, r.get("error", "移除失败"))
    return r


# ================================================================ 目录树 / 读 / 写
def tree(token: str, subdir: str = "") -> dict:
    """返回某一层的直接子项（前端逐级展开，懒加载）。"""
    aid, un, wd, _ = _ctx(token)
    try:
        root = guard.resolve_in_workdir(subdir, wd) if (subdir or "").strip() else wd
    except ValueError as e:
        raise HTTPException(400, str(e))
    if not root.is_dir():
        raise HTTPException(400, "%s 不是目录" % (subdir or "."))
    entries = []
    try:
        children = sorted(root.iterdir(), key=lambda p: (not p.is_dir(), p.name.lower()))
    except OSError as e:
        raise HTTPException(400, "目录读取失败：%s" % e)
    for p in children:
        if p.name == guard.RECYCLE_DIR:
            is_dir, size = True, 0                       # 回收站显示，但内容不在这里列
        else:
            is_dir = p.is_dir()
            try:
                size = 0 if is_dir else p.stat().st_size
            except OSError:
                size = 0
        try:
            mtime = p.stat().st_mtime
        except OSError:
            mtime = 0
        entries.append({
            "name": p.name, "rel": _rel(p, wd), "is_dir": is_dir, "size": size, "mtime": mtime,
            "credential": guard.is_credential(p, wd),
            "protected": bool(guard.is_forbidden_dir(p)) or p.name == guard.PROTECT_FILE
                         or bool(_protected_quietly(p, wd)),
        })
    parent = "" if root == Path(wd) else _rel(root.parent, wd)
    return {"ok": True, "path": _rel(root, wd), "parent": parent, "entries": entries,
            "is_root": root == Path(wd)}


def _protected_quietly(p: Path, wd: Path) -> str | None:
    """保护名单查询，但**不因名单缺失而报错**（列表用；写入侧仍 fail-closed）。"""
    try:
        return guard.protected_reason(p, wd)
    except Exception:                              # noqa: BLE001
        return None


def read_(token: str, path: str) -> dict:
    aid, un, wd, _ = _ctx(token)
    try:
        p = guard.resolve_in_workdir(path, wd)
    except ValueError as e:
        raise HTTPException(400, str(e))
    reason = guard.check_readable(p, wd)
    if reason:
        _audit("read", str(p), False, reason, aid, wd)
        raise HTTPException(403, reason)
    if not p.is_file():
        raise HTTPException(404, "%s 不是文件" % path)
    size = p.stat().st_size
    if size > _MAX_READ * 4:
        raise HTTPException(400, "文件 %s 过大（%s），请下载查看" % (path, _human(size)))
    ext = p.suffix.lower()
    kind, text = "download", ""
    if ext in _TEXT_EXT:
        kind = "text"
        try:
            text = p.read_text(encoding="utf-8", errors="replace")
        except OSError as e:
            raise HTTPException(400, "读取失败：%s" % e)
    elif ext == ".docx":
        kind, text = "preview", _docx_text(p)
    elif ext == ".xlsx":
        kind, text = "preview", _xlsx_text(p)
    if len(text) > _MAX_READ:
        text = text[:_MAX_READ] + "\n… 预览已截断"
    _audit("read", str(p), True, "%s/%s" % (kind, size), aid, wd)
    return {"ok": True, "kind": kind, "text": text, "name": p.name, "rel": _rel(p, wd),
            "size": size, "mtime": p.stat().st_mtime,           # 亚秒：前端原样回传做冲突检测
            "editable": ext in _EDIT_EXT and not guard.is_credential(p, wd),
            "protected": p.name == guard.PROTECT_FILE or bool(_protected_quietly(p, wd))}


def download(token: str, path: str):
    """仅下载（PDF 等不可预览类型；也允许任意非凭据文件下载）。"""
    aid, un, wd, _ = _ctx(token)
    try:
        p = guard.resolve_in_workdir(path, wd)
    except ValueError as e:
        raise HTTPException(400, str(e))
    reason = guard.check_readable(p, wd)
    if reason:
        raise HTTPException(403, reason)
    if not p.is_file():
        raise HTTPException(404, "%s 不是文件" % path)
    _audit("download", str(p), True, str(p.stat().st_size), aid, wd)
    return FileResponse(str(p), filename=p.name)


def write_(token: str, req: WriteReq) -> dict:
    aid, un, wd, _ = _ctx(token)
    try:
        p = guard.resolve_in_workdir(req.path, wd)
    except ValueError as e:
        raise HTTPException(400, str(e))
    reason = guard.assert_writable(p, wd)
    if reason:
        _audit("write", str(p), False, reason, aid, wd)
        raise HTTPException(403, reason)
    if not p.is_file():
        raise HTTPException(404, "%s 不存在（新建请用 create）" % req.path)
    if req.mtime is not None:
        # ★ 亚秒精度：整秒会让"同一秒内被改写"漏检（单测实测踩到）
        cur = p.stat().st_mtime
        if abs(cur - float(req.mtime)) > 1e-6:
            _audit("write", str(p), False, "mtime 冲突", aid, wd)
            raise HTTPException(409, "文件已被改动（可能是 Agent 或别的窗口写的），请重新打开再改。")
    blob = (req.content or "").encode("utf-8")
    if len(blob) > _max_write():
        raise HTTPException(400, "内容 %s 超过单次写入上限 %s" % (_human(len(blob)), _human(_max_write())))
    try:
        p.write_bytes(blob)
    except OSError as e:
        raise HTTPException(400, "保存失败：%s" % e)
    _audit("write", str(p), True, _human(len(blob)), aid, wd)
    return {"ok": True, "mtime": p.stat().st_mtime, "size": len(blob)}


def create_(token: str, req: CreateReq) -> dict:
    aid, un, wd, _ = _ctx(token)
    try:
        p = guard.resolve_in_workdir(req.path, wd)
    except ValueError as e:
        raise HTTPException(400, str(e))
    reason = guard.assert_writable(p, wd)
    if reason:
        _audit("create", str(p), False, reason, aid, wd)
        raise HTTPException(403, reason)
    if p.exists():
        raise HTTPException(409, "%s 已存在" % req.path)
    blob = (req.content or "").encode("utf-8")
    if len(blob) > _max_write():
        raise HTTPException(400, "内容超过单次写入上限 %s" % _human(_max_write()))
    try:
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_bytes(blob)
    except OSError as e:
        raise HTTPException(400, "创建失败：%s" % e)
    _audit("create", str(p), True, _human(len(blob)), aid, wd)
    return {"ok": True, "rel": _rel(p, wd)}


def delete_(token: str, req: PathReq) -> dict:
    aid, un, wd, _ = _ctx(token)
    try:
        p = guard.resolve_in_workdir(req.path, wd)
    except ValueError as e:
        raise HTTPException(400, str(e))
    if p == Path(wd).resolve():
        raise HTTPException(400, "不能删除工作目录根目录本身")
    if not p.exists():
        raise HTTPException(404, "%s 不存在" % req.path)
    reason = guard.assert_writable(p, wd)
    if reason:
        _audit("delete", str(p), False, reason, aid, wd)
        raise HTTPException(403, reason)
    try:
        dst = guard.recycle(wd, p)
    except OSError as e:
        raise HTTPException(400, "删除失败：%s" % e)
    try:
        guard.purge_expired_trash(wd, _ttl_days())          # 每次删除顺带清过期项
    except Exception:                              # noqa: BLE001
        pass
    _audit("delete", str(p), True, "→ %s" % dst.name, aid, wd)
    return {"ok": True, "trash_name": dst.name}


def rename_(token: str, req: RenameReq) -> dict:
    aid, un, wd, _ = _ctx(token)
    try:
        src = guard.resolve_in_workdir(req.path, wd)
        dst = guard.resolve_in_workdir(req.target, wd)
    except ValueError as e:
        raise HTTPException(400, str(e))
    # ★ 门禁先于存在性检查：受保护目标要报"受保护"这个真实原因，也避免用报错差异探测文件是否存在
    for p in (src, dst):
        reason = guard.assert_writable(p, wd)
        if reason:
            _audit("rename", "%s → %s" % (src, dst), False, reason, aid, wd)
            raise HTTPException(403, reason)
    if not src.exists():
        raise HTTPException(404, "%s 不存在" % req.path)
    if dst.exists():
        raise HTTPException(409, "%s 已存在" % req.target)
    try:
        dst.parent.mkdir(parents=True, exist_ok=True)
        src.rename(dst)
    except OSError as e:
        raise HTTPException(400, "重命名失败：%s" % e)
    _audit("rename", "%s → %s" % (_rel(src, wd), _rel(dst, wd)), True, "", aid, wd)
    return {"ok": True, "rel": _rel(dst, wd)}


# ================================================================ 回收站
def trash_(token: str, req: TrashReq) -> dict:
    aid, un, wd, _ = _ctx(token)
    act = (req.action or "list").strip().lower()
    if act == "list":
        return {"ok": True, "items": guard.trash_entries(wd), "ttl_days": _ttl_days()}
    if act == "restore":
        if not req.name:
            raise HTTPException(400, "restore 需要 name")
        try:
            dst = guard.trash_restore(wd, req.name, req.target)
        except (FileNotFoundError, ValueError, OSError) as e:
            raise HTTPException(400, "还原失败：%s" % e)
        _audit("trash_restore", req.name, True, "→ %s" % _rel(dst, wd), aid, wd)
        return {"ok": True, "rel": _rel(dst, wd)}
    if act == "purge":
        n = guard.trash_purge(wd, req.name)
        _audit("trash_purge", req.name or "(清空)", True, "%d 项" % n, aid, wd)
        return {"ok": True, "purged": n}
    raise HTTPException(400, "action 只能是 list / restore / purge")


# ================================================================ 预览提取
def _docx_text(p: Path) -> str:
    try:
        import docx
    except ImportError:
        return "（本机没有 python-docx，无法预览，请下载查看）"
    try:
        d = docx.Document(str(p))
    except Exception as e:                         # noqa: BLE001
        return "（.docx 解析失败：%s，请下载查看）" % e
    out = [x.text for x in d.paragraphs if x.text.strip()]
    for t in d.tables:
        out.append("—— 表格 ——")
        for row in t.rows:
            out.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(out) or "（空文档）"


def _xlsx_text(p: Path, max_rows: int = 200) -> str:
    try:
        import openpyxl
    except ImportError:
        return "（本机没有 openpyxl，无法预览，请下载查看）"
    try:
        wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
    except Exception as e:                         # noqa: BLE001
        return "（.xlsx 解析失败：%s，请下载查看）" % e
    out = []
    for ws in wb.worksheets:
        out.append("—— 工作表：%s（%s 行 × %s 列）——" % (ws.title, ws.max_row, ws.max_column))
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                out.append("… 只预览前 %d 行" % max_rows)
                break
            out.append(" | ".join("" if c is None else str(c) for c in row))
    try:
        wb.close()
    except Exception:                              # noqa: BLE001
        pass
    return "\n".join(out) or "（空表）"
